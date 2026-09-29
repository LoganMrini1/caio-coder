"""Turn the raw postings spreadsheet into a slim data/postings.csv.

Usage (from the project root):
    python src/make_slim.py data/Manual_CAIO_Records.xlsx
Keeps only: url, company, title, location, post_date, description (HTML stripped to plain text).
"""
import csv
import html
import re
import sys
from html.parser import HTMLParser

from openpyxl import load_workbook

KEEP = {  # output name -> header text in the raw file
    "url": "Posting URL",
    "company": "Company",
    "title": "Job Title",
    "location": "Location",
    "post_date": "Post Date",
    "description": "Description",
}
BLOCK_TAGS = {"p", "div", "br", "li", "ul", "ol", "h1", "h2", "h3", "h4", "h5", "h6", "tr"}


class _Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.out = []

    def handle_starttag(self, tag, attrs):
        if tag == "li":
            self.out.append("\n- ")
        elif tag in BLOCK_TAGS:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in BLOCK_TAGS:
            self.out.append("\n")

    def handle_data(self, data):
        self.out.append(data)


def html_to_text(raw):
    if not raw:
        return ""
    p = _Text()
    p.feed(str(raw))
    text = html.unescape("".join(p.out)).replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def main(path, out_path="data/postings.csv"):
    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    # header row = first row that contains "Posting URL" (raw files can have blank rows above it)
    hi = next(i for i, r in enumerate(rows) if r and "Posting URL" in r)
    idx = {name: rows[hi].index(src) for name, src in KEEP.items()}

    seen, kept, no_desc, dupes = set(), [], 0, 0
    for r in rows[hi + 1:]:
        url = (r[idx["url"]] or "").strip() if r[idx["url"]] else ""
        if not url:
            continue
        if url in seen:
            dupes += 1
            continue
        seen.add(url)
        rec = {k: ("" if r[i] is None else str(r[i]).strip()) for k, i in idx.items()}
        rec["description"] = html_to_text(r[idx["description"]])
        if len(rec["description"]) < 200:
            no_desc += 1
        kept.append(rec)

    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(KEEP))
        w.writeheader()
        w.writerows(kept)
    print(f"Wrote {len(kept)} postings to {out_path}")
    print(f"Skipped duplicate URLs: {dupes} | Short/empty descriptions (<200 chars): {no_desc}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("Usage: python src/make_slim.py <raw_xlsx_path>")
    main(sys.argv[1])