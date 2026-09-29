"""Maps workbook columns (sheet CAIO Data) to fields and handles formula copying.

Observed in the provided workbook:
- Formulas exist ONLY in row 2 (the example row). New rows need those formulas copied down.
- Data-validation dropdowns cover rows 2:202 only.
"""
from openpyxl.formula.translate import Translator

SHEET = "CAIO Data"
FIRST_DATA_ROW = 3          # row 2 is the example row

# Formula / helper columns: never write values here, copy the formula from row 2
FORMULA_COLS = ["A", "AI", "AO", "AP", "AR", "BB", "BC", "BG", "BH", "BI"]

# Columns filled with constants by code (not the LLM)
CONST_COLS = {"date": "B", "collected_by": "C", "source_category": "D",
              "evidence_type": "U", "source_url": "V", "pdf": "W"}

FIELD_COLS = {
    "organization_name": "E", "sector": "F", "industry": "G", "org_size": "H",
    "headquarters": "I", "exact_title": "J", "title_bucket": "K",
    "employment_type": "M", "reporting_line": "R", "scope": "S", "team_structure": "T",
    "notes": "BF",
}
DIM_COLS = {"strategic": "Y", "technical": "Z", "data": "AA", "cybersecurity": "AB",
            "risk_compliance": "AC", "stakeholder": "AD", "governance": "AE",
            "workforce": "AF", "vendor_ecosystem": "AG", "value": "AH"}
DR_COLS = {"low": "AK", "moderate": "AL", "stronger": "AM", "high": "AN"}
DR_QUOTE_COL = "AJ"
CYBER_COL = "AS"
INTERACT_COLS = {"cio": "AT", "ciso": "AU", "legal": "AV", "business_units": "AW",
                 "hr": "AX", "procurement": "AY", "privacy": "AZ", "board": "BA"}
QUOTE_NOTES_COL = "X"       # evidence quotes go here
MATURITY_COL, ARCHETYPE_COL = "BD", "BE"   # derived in code (derive.py, next step)


def copy_formulas(ws, target_row, template_row=2):
    for col in FORMULA_COLS:
        src = ws[f"{col}{template_row}"].value
        if isinstance(src, str) and src.startswith("="):
            ws[f"{col}{target_row}"] = Translator(src, origin=f"{col}{template_row}").translate_formula(f"{col}{target_row}")


def next_empty_row(ws):
    # A row counts as used if it has a source URL (company name can legitimately be blank)
    col = CONST_COLS["source_url"]
    r = FIRST_DATA_ROW
    while ws[f"{col}{r}"].value not in (None, ""):
        r += 1
    return r