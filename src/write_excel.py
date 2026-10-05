"""
Create and write the CAIO coding results workbook.

This version creates a fresh workbook instead of modifying the original
template in place. It provides:
- 61 CAIO Data columns (A:BI)
- 1,000 data rows
- dropdown validations
- formulas for derived fields
- conditional formatting
- frozen headers
- Excel table
- hidden Lists sheet
- compatibility with the existing run.py pipeline
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

try:
    from openpyxl.workbook.properties import CalcProperties
except ImportError:
    CalcProperties = None

from .schema import PostingCoding


MAX_DATA_ROWS = 1000
FIRST_DATA_ROW = 3
LAST_DATA_ROW = FIRST_DATA_ROW + MAX_DATA_ROWS - 1


HEADERS = [
    "Record ID",
    "Date Collected",
    "Collected By",
    "Collection Source Category",
    "Organization Name",
    "Sector",
    "Industry",
    "Org Size",
    "Headquarters",
    "Exact Title",
    "CAIO or Adjacent Title",
    "Appointment Date",
    "Employment Type",
    "Person Name",
    "Prior Role",
    "Prior Function",
    "Background Type",
    "Reporting Line",
    "Scope",
    "Team Structure",
    "Source Evidence Type",
    "Source URL",
    "PDF of Posting",
    "Source Quote / Notes",
    "Strategic Orchestration Score",
    "Technical Orchestration Score",
    "Data Orchestration Score",
    "Cybersecurity Orchestration Score",
    "Risk/Compliance Orchestration Score",
    "Stakeholder Orchestration Score",
    "Governance Orchestration Score",
    "Workforce Orchestration Score",
    "Vendor/Ecosystem Orchestration Score",
    "Value Orchestration Score",
    "Orchestration Breadth Score (0-30)",
    "Decision-Rights Language Observed",
    "DR Language: Low (advise/support/recommend)",
    "DR Language: Moderate (coordinate/partner/facilitate)",
    "DR Language: Stronger (lead/oversee/chair/govern)",
    "DR Language: High (approve/prioritize/policy/enforce/halt/waive)",
    "Decision-Rights Language Breadth (0-4)",
    "Decision-Rights Strength (auto)",
    "Decision-Rights Strength (Override — optional)",
    "Decision-Rights Strength (Final)",
    "Cybersecurity Aspect Score",
    "CIO",
    "CISO",
    "Legal",
    "Business Units",
    "HR",
    "Procurement",
    "Privacy",
    "Board",
    "Interacts Span Count",
    "Suggested Maturity Level (auto)",
    "Orchestration Maturity Level",
    "CAIO Archetype",
    "General Notes / Flags",
    "Helper: Decision Rank",
    "Helper: Base Level",
    "Helper: Suggested Level #",
]


FORMULA_COLUMNS = {
    "AI",
    "AO",
    "AP",
    "AR",
    "BB",
    "BC",
    "BG",
    "BH",
    "BI",
}


LISTS = {
    "Initials": [
        "Logan",
        "Roman",
    ],
    "SourceCategory": [
        "Current CAIO/Equivalent",
        "Job Posting",
        "WRDS",
    ],
    "Sector": [
        "Public",
        "Private",
        "Government",
        "Nonprofit",
    ],
    "TitleBucket": [
        "CAIO",
        "VP of AI",
        "CDAO",
        "AI Leader",
        "Other",
    ],
    "EmploymentType": [
        "Full-time",
        "Interim",
        "Fractional",
        "Unknown",
    ],
    "BackgroundType": [
        "Technical",
        "Business",
        "Legal",
        "Cyber",
        "Mixed",
        "Unknown",
    ],
    "ReportingLine": [
        "CEO",
        "CIO",
        "CTO",
        "CDO",
        "CISO",
        "COO",
        "Legal",
        "Risk",
        "Board",
        "Unknown",
    ],
    "Scope": [
        "Enterprise-wide",
        "Business-unit",
        "Product",
        "Data/Analytics",
        "Governance",
        "Federal Agency",
        "Regional",
    ],
    "TeamStructure": [
        "Owns AI team",
        "Chairs AI council",
        "Embedded in IT",
        "Works through matrix structure",
        "Unknown",
    ],
    "SourceEvidenceType": [
        "Press release",
        "Job posting",
        "LinkedIn",
        "Annual report",
        "Agency page",
        "Executive bio",
        "Other social media (X)",
        "Other",
    ],
    "Score0to3": [
        0,
        1,
        2,
        3,
    ],
    "DecisionRights": [
        "Yes",
        "No",
    ],
    "YesNo": [
        "Yes",
        "No",
    ],
    "MaturityLevel": [
        "Level 1: Symbolic CAIO",
        "Level 2: Advisory CAIO",
        "Level 3: Coordinating CAIO",
        "Level 4: Governing CAIO",
        "Level 5: Enterprise Orchestrator",
    ],
    "Archetype": [
        "AI Strategy Orchestrator",
        "AI Governance Orchestrator",
        "AI Technology Orchestrator",
        "Cyber/IS Orchestrator",
        "Workforce Transformation Orchestrator",
        "Hybrid Enterprise Orchestrator",
    ],
}


def _safe(value):
    """Convert empty strings to None while preserving useful values."""
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def _evidence_notes(coding: PostingCoding) -> str:
    """Build a compact evidence string from the coded dimensions."""

    parts = []

    dimensions = [
        ("strategic", "Strategic"),
        ("technical", "Technical"),
        ("data", "Data"),
        ("cybersecurity", "Cybersecurity"),
        ("risk_compliance", "Risk/Compliance"),
        ("stakeholder", "Stakeholder"),
        ("governance", "Governance"),
        ("workforce", "Workforce"),
        ("vendor_ecosystem", "Vendor/Ecosystem"),
        ("value", "Value"),
    ]

    for name, label in dimensions:
        dimension = getattr(coding.dimensions, name)

        if dimension.score > 0 and dimension.evidence:
            parts.append(
                f"{label}: {dimension.evidence}"
            )

    if (
        coding.cyber_score.score > 0
        and coding.cyber_score.evidence
    ):
        parts.append(
            f"Cybersecurity Aspect: {coding.cyber_score.evidence}"
        )

    return " | ".join(parts)


def _build_values(
    posting: dict,
    coding: PostingCoding,
    collected_by: str,
):
    """Build non-formula values for one result row."""

    values = {
        "A": _safe(posting.get("record_id")),
        "B": _safe(posting.get("date_collected")),
        "C": collected_by,
        "D": _safe(posting.get("source_category")),
        "E": _safe(coding.organization_name),
        "F": _safe(coding.sector),
        "G": _safe(coding.industry),
        "H": _safe(coding.org_size),
        "I": _safe(coding.headquarters),
        "J": _safe(coding.exact_title),
        "K": _safe(coding.title_bucket),
        "L": _safe(coding.appointment_date),
        "M": _safe(coding.employment_type),
        "N": _safe(coding.person_name),
        "O": _safe(coding.prior_role),
        "P": _safe(coding.prior_function),
        "Q": _safe(coding.background_type),
        "R": _safe(coding.reporting_line),
        "S": _safe(coding.scope),
        "T": _safe(coding.team_structure),
        "U": _safe(posting.get("source_evidence_type")),
        "V": _safe(posting.get("source_url")),
        "W": _safe(posting.get("pdf_of_posting")),
        "X": _evidence_notes(coding),

        "Y": coding.dimensions.strategic.score,
        "Z": coding.dimensions.technical.score,
        "AA": coding.dimensions.data.score,
        "AB": coding.dimensions.cybersecurity.score,
        "AC": coding.dimensions.risk_compliance.score,
        "AD": coding.dimensions.stakeholder.score,
        "AE": coding.dimensions.governance.score,
        "AF": coding.dimensions.workforce.score,
        "AG": coding.dimensions.vendor_ecosystem.score,
        "AH": coding.dimensions.value.score,

        "AJ": coding.decision_rights.observed_language,

        "AK": (
            "Yes"
            if coding.decision_rights.low
            else "No"
        ),
        "AL": (
            "Yes"
            if coding.decision_rights.moderate
            else "No"
        ),
        "AM": (
            "Yes"
            if coding.decision_rights.stronger
            else "No"
        ),
        "AN": (
            "Yes"
            if coding.decision_rights.high
            else "No"
        ),

        "AQ": _safe(
            coding.decision_rights_override
        ),

        "AS": coding.cyber_score.score,

        "AT": (
            "Yes"
            if coding.interacts.cio
            else "No"
        ),
        "AU": (
            "Yes"
            if coding.interacts.ciso
            else "No"
        ),
        "AV": (
            "Yes"
            if coding.interacts.legal
            else "No"
        ),
        "AW": (
            "Yes"
            if coding.interacts.business_units
            else "No"
        ),
        "AX": (
            "Yes"
            if coding.interacts.hr
            else "No"
        ),
        "AY": (
            "Yes"
            if coding.interacts.procurement
            else "No"
        ),
        "AZ": (
            "Yes"
            if coding.interacts.privacy
            else "No"
        ),
        "BA": (
            "Yes"
            if coding.interacts.board
            else "No"
        ),

        "BD": coding.orchestration_maturity_level,
        "BE": coding.caio_archetype,
        "BF": coding.general_notes,
    }

    if coding.interacts.other:
        if values["BF"]:
            values["BF"] += (
                f" | Other interactions: "
                f"{coding.interacts.other}"
            )
        else:
            values["BF"] = (
                f"Other interactions: "
                f"{coding.interacts.other}"
            )

    if coding.source_quote_notes:
        if values["BF"]:
            values["BF"] += (
                f" | Source notes: "
                f"{coding.source_quote_notes}"
            )
        else:
            values["BF"] = coding.source_quote_notes

    return values


def _formula(column: str, row: int) -> str:
    """Return the Excel formula for a calculated column."""

    if column == "AI":
        return f"=SUM(Y{row}:AH{row})"

    if column == "AO":
        return f'=COUNTIF(AK{row}:AN{row},"Yes")'

    if column == "AP":
        return (
            f'=IF(AN{row}="Yes",'
            f'"High decision-rights authority",'
            f'IF(AM{row}="Yes",'
            f'"Stronger authority",'
            f'IF(AL{row}="Yes",'
            f'"Orchestration role, moderate authority",'
            f'IF(AK{row}="Yes",'
            f'"Low direct authority",""))))'
        )

    if column == "AR":
        return (
            f'=IF(AQ{row}<>"",AQ{row},AP{row})'
        )

    if column == "BB":
        return (
            f'=COUNTIF(AT{row}:BA{row},"Yes")'
        )

    if column == "BC":
        return (
            f'=IF(BI{row}=1,'
            f'"Level 1: Symbolic CAIO",'
            f'IF(BI{row}=2,'
            f'"Level 2: Advisory CAIO",'
            f'IF(BI{row}=3,'
            f'"Level 3: Coordinating CAIO",'
            f'IF(BI{row}=4,'
            f'"Level 4: Governing CAIO",'
            f'IF(BI{row}=5,'
            f'"Level 5: Enterprise Orchestrator","")))))'
        )

    if column == "BG":
        return (
            f'=IF(AN{row}="Yes",4,'
            f'IF(AM{row}="Yes",3,'
            f'IF(AL{row}="Yes",2,'
            f'IF(AK{row}="Yes",1,0))))'
        )

    if column == "BH":
        return (
            f'=IF(AI{row}<=5,1,'
            f'IF(AI{row}<=11,2,'
            f'IF(AI{row}<=17,3,'
            f'IF(AI{row}<=23,4,5))))'
        )

    if column == "BI":
        return (
            f'=IF(BH{row}>=4,'
            f'IF(BG{row}<3,3,'
            f'IF(BH{row}=5,'
            f'IF(AND(BB{row}>=6,AN{row}="Yes"),5,4),'
            f'4)),'
            f'BH{row})'
        )

    return ""


def _add_validations(ws):
    """Add dropdown validations to the data-entry ranges."""

    def add_list_validation(
        formula,
        cell_range,
        prompt_title=None,
        prompt=None,
    ):
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
        )

        if prompt_title:
            dv.promptTitle = prompt_title

        if prompt:
            dv.prompt = prompt

        dv.errorTitle = "Invalid value"
        dv.error = (
            "Please select a value from the dropdown."
        )
        dv.showErrorMessage = True
        dv.showInputMessage = True

        ws.add_data_validation(dv)
        dv.add(cell_range)

    # Collected By
    add_list_validation(
        "=Lists!$A$2:$A$3",
        f"C{FIRST_DATA_ROW}:C{LAST_DATA_ROW}",
    )

    # Collection Source Category
    add_list_validation(
        "=Lists!$B$2:$B$4",
        f"D{FIRST_DATA_ROW}:D{LAST_DATA_ROW}",
    )

    # Sector
    add_list_validation(
        "=Lists!$C$2:$C$5",
        f"F{FIRST_DATA_ROW}:F{LAST_DATA_ROW}",
    )

    # CAIO / Adjacent Title
    add_list_validation(
        "=Lists!$D$2:$D$6",
        f"K{FIRST_DATA_ROW}:K{LAST_DATA_ROW}",
    )

    # Employment Type
    add_list_validation(
        "=Lists!$E$2:$E$5",
        f"M{FIRST_DATA_ROW}:M{LAST_DATA_ROW}",
    )

    # Background Type
    add_list_validation(
        "=Lists!$F$2:$F$7",
        f"Q{FIRST_DATA_ROW}:Q{LAST_DATA_ROW}",
    )

    # Reporting Line
    add_list_validation(
        "=Lists!$G$2:$G$11",
        f"R{FIRST_DATA_ROW}:R{LAST_DATA_ROW}",
    )

    # Scope
    add_list_validation(
        "=Lists!$H$2:$H$8",
        f"S{FIRST_DATA_ROW}:S{LAST_DATA_ROW}",
    )

    # Team Structure
    add_list_validation(
        "=Lists!$I$2:$I$6",
        f"T{FIRST_DATA_ROW}:T{LAST_DATA_ROW}",
    )

    # Source Evidence Type
    add_list_validation(
        "=Lists!$J$2:$J$9",
        f"U{FIRST_DATA_ROW}:U{LAST_DATA_ROW}",
    )

    # Ten orchestration dimensions: 0-3
    add_list_validation(
        "=Lists!$K$2:$K$5",
        f"Y{FIRST_DATA_ROW}:AH{LAST_DATA_ROW}",
    )

    # Decision-rights language: Yes/No
    add_list_validation(
        "=Lists!$M$2:$M$3",
        f"AK{FIRST_DATA_ROW}:AN{LAST_DATA_ROW}",
    )

    # Cybersecurity score: 0-3
    add_list_validation(
        "=Lists!$K$2:$K$5",
        f"AS{FIRST_DATA_ROW}:AS{LAST_DATA_ROW}",
    )

    # Interacts-with fields: Yes/No
    add_list_validation(
        "=Lists!$M$2:$M$3",
        f"AT{FIRST_DATA_ROW}:BA{LAST_DATA_ROW}",
    )

    # Manual maturity level
    add_list_validation(
        "=Lists!$N$2:$N$6",
        f"BD{FIRST_DATA_ROW}:BD{LAST_DATA_ROW}",
    )

    # Archetype
    add_list_validation(
        "=Lists!$O$2:$O$7",
        f"BE{FIRST_DATA_ROW}:BE{LAST_DATA_ROW}",
    )


def _style_sheet(ws):
    """Apply workbook formatting."""

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    thin_gray = Side(
        style="thin",
        color="D9E1F2",
    )

    header_border = Border(
        bottom=thin_gray,
    )

    for cell in ws[2]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = header_border

    ws.row_dimensions[2].height = 60

    # Title.
    ws["A1"] = (
        "CAIO / AI Leadership Job Posting Coding Results"
    )

    ws["A1"].font = Font(
        bold=True,
        size=14,
    )

    ws["A1"].alignment = Alignment(
        vertical="center",
    )

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(HEADERS),
    )

    ws.row_dimensions[1].height = 26

    # General cell formatting.
    for row in ws.iter_rows(
        min_row=FIRST_DATA_ROW,
        max_row=LAST_DATA_ROW,
        min_col=1,
        max_col=len(HEADERS),
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    # Helper columns.
    helper_fill = PatternFill(
        fill_type="solid",
        fgColor="E7E6E6",
    )

    for column in FORMULA_COLUMNS:
        column_number = _column_number(column)

        for row in range(
            FIRST_DATA_ROW,
            LAST_DATA_ROW + 1,
        ):
            ws.cell(
                row=row,
                column=column_number,
            ).fill = helper_fill

    # Date formatting.
    for row in range(
        FIRST_DATA_ROW,
        LAST_DATA_ROW + 1,
    ):
        ws[f"B{row}"].number_format = "m/d/yyyy"

    widths = {
        "A": 12,
        "B": 14,
        "C": 14,
        "D": 24,
        "E": 28,
        "F": 14,
        "G": 24,
        "H": 18,
        "I": 24,
        "J": 34,
        "K": 18,
        "L": 16,
        "M": 16,
        "N": 24,
        "O": 28,
        "P": 22,
        "Q": 18,
        "R": 16,
        "S": 20,
        "T": 28,
        "U": 20,
        "V": 45,
        "W": 30,
        "X": 70,
        "Y": 14,
        "Z": 14,
        "AA": 14,
        "AB": 16,
        "AC": 18,
        "AD": 16,
        "AE": 16,
        "AF": 16,
        "AG": 18,
        "AH": 14,
        "AI": 18,
        "AJ": 35,
        "AK": 18,
        "AL": 20,
        "AM": 20,
        "AN": 20,
        "AO": 18,
        "AP": 30,
        "AQ": 32,
        "AR": 30,
        "AS": 18,
        "AT": 12,
        "AU": 12,
        "AV": 12,
        "AW": 16,
        "AX": 12,
        "AY": 16,
        "AZ": 12,
        "BA": 12,
        "BB": 16,
        "BC": 28,
        "BD": 28,
        "BE": 30,
        "BF": 60,
        "BG": 18,
        "BH": 16,
        "BI": 18,
    }

    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.auto_filter.ref = (
        f"A2:BI{LAST_DATA_ROW}"
    )

    # Maturity conditional formatting.
    ws.conditional_formatting.add(
        f"BC{FIRST_DATA_ROW}:BC{LAST_DATA_ROW}",
        FormulaRule(
            formula=[
                f'LEFT(BC{FIRST_DATA_ROW},7)="Level 5"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor="C6E0B4",
            ),
        ),
    )

    ws.conditional_formatting.add(
        f"BC{FIRST_DATA_ROW}:BC{LAST_DATA_ROW}",
        FormulaRule(
            formula=[
                f'LEFT(BC{FIRST_DATA_ROW},7)="Level 4"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor="DDEBF7",
            ),
        ),
    )

    ws.conditional_formatting.add(
        f"BC{FIRST_DATA_ROW}:BC{LAST_DATA_ROW}",
        FormulaRule(
            formula=[
                f'LEFT(BC{FIRST_DATA_ROW},7)="Level 3"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor="FFF2CC",
            ),
        ),
    )

    ws.conditional_formatting.add(
        f"BC{FIRST_DATA_ROW}:BC{LAST_DATA_ROW}",
        FormulaRule(
            formula=[
                f'LEFT(BC{FIRST_DATA_ROW},7)="Level 1"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor="F4CCCC",
            ),
        ),
    )


def _column_number(column_letter: str) -> int:
    """Convert an Excel column letter to a numeric index."""

    result = 0

    for char in column_letter.upper():
        result = (
            result * 26
            + (ord(char) - ord("A") + 1)
        )

    return result


def _build_lists_sheet(wb):
    """Create the hidden Lists worksheet."""

    ws = wb.create_sheet("Lists")

    list_names = list(LISTS.keys())

    for col_index, list_name in enumerate(
        list_names,
        start=1,
    ):
        ws.cell(
            row=1,
            column=col_index,
            value=list_name,
        )

        ws.cell(
            row=1,
            column=col_index,
        ).font = Font(
            bold=True
        )

        for row_index, value in enumerate(
            LISTS[list_name],
            start=2,
        ):
            ws.cell(
                row=row_index,
                column=col_index,
                value=value,
            )

        ws.column_dimensions[
            get_column_letter(col_index)
        ].width = max(
            15,
            min(45, len(list_name) + 5),
        )

    ws.sheet_state = "hidden"

    return ws


def prepare_output(template_path, output_path):
    """
    Create a fresh CAIO results workbook.

    template_path remains in the function signature because run.py passes it.
    The old template is intentionally not copied or edited.
    """

    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    print("Creating fresh output workbook...")

    wb = Workbook()

    ws = wb.active
    ws.title = "CAIO Data"

    # Row 1: title.
    ws.append(
        [
            "CAIO / AI Leadership Job Posting Coding Results"
        ]
    )

    # Row 2: headers.
    ws.append(HEADERS)

    # Pre-create 1,000 data rows and formulas.
    for row in range(
        FIRST_DATA_ROW,
        LAST_DATA_ROW + 1,
    ):
        for column in FORMULA_COLUMNS:
            column_number = _column_number(column)

            ws.cell(
                row=row,
                column=column_number,
                value=_formula(column, row),
            )

    _build_lists_sheet(wb)

    _add_validations(ws)

    _style_sheet(ws)

    table = Table(
        displayName="CAIOData",
        ref=f"A2:BI{LAST_DATA_ROW}",
    )

    table_style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    table.tableStyleInfo = table_style

    ws.add_table(table)

    if CalcProperties is not None:
        try:
            wb.calculation = CalcProperties(
                calcMode="auto",
                fullCalcOnLoad=True,
                forceFullCalc=True,
            )
        except Exception:
            pass

    wb.save(output_path)

    print(
        f"Prepared workbook through Excel row {LAST_DATA_ROW}."
    )


def write_posting(
    output_path,
    posting,
    coding,
    row_num,
    collected_by,
):
    """
    Write one coded posting to the workbook.

    IMPORTANT:
    row_num intentionally matches the keyword used by run.py.
    """

    output_path = Path(output_path)

    wb = load_workbook(
        output_path,
        data_only=False,
    )

    ws = wb["CAIO Data"]

    values = _build_values(
        posting=posting,
        coding=coding,
        collected_by=collected_by,
    )

    # Write all non-formula values.
    for column, value in values.items():
        if column in FORMULA_COLUMNS:
            continue

        ws[f"{column}{row_num}"] = value

    # Reassert formulas.
    for column in FORMULA_COLUMNS:
        ws[f"{column}{row_num}"] = _formula(
            column,
            row_num,
        )

    # Date format.
    ws[f"B{row_num}"].number_format = "m/d/yyyy"

    # Wrap text.
    for column_number in range(
        1,
        len(HEADERS) + 1,
    ):
        ws.cell(
            row=row_num,
            column=column_number,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    # Preserve helper-column shading.
    helper_fill = PatternFill(
        fill_type="solid",
        fgColor="E7E6E6",
    )

    for column in FORMULA_COLUMNS:
        ws[f"{column}{row_num}"].fill = helper_fill

    wb.save(output_path)